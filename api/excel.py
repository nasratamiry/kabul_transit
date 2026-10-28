"""
Excel (.xlsx) exports: the income statement and the transactions list.

Right-to-left sheets, Dari headers, Vazirmatn font (Excel falls back to a system font
when it is not installed), real numbers with two decimals (never text), Shamsi and
Gregorian date columns, a title block (branch, period, generation time), a frozen header
row and set column widths. Each builder returns the file's bytes; xlsx_response() serves them.
"""
from decimal import Decimal
from io import BytesIO

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from finance.models import StatementLine

FONT = 'Vazirmatn'
BRAND = '0C514C'
BRAND_LIGHT = 'E6F0EF'
MONEY = '#,##0.00;[Red](#,##0.00)'
RATE = '#,##0.00######'
THIN = Side(style='thin', color='D1D5DB')

SUBTOTALS = {  # after the lines of this section: (label, totals key)
    StatementLine.Section.COGS: ('مفاد ناخالص', 'gross_profit'),
}
SECTION_TOTAL_KEY = {'REVENUE': 'revenue', 'COGS': 'cogs', 'OPEX': 'opex', 'OTHER': 'other',
                     'OTHER_INCOME': 'other_income'}


def _ltr(text):
    """A date (or date + time) inside RTL text: left-to-right marks keep '1405-08-01' from being reversed."""
    return f'\u200e{text}\u200e'


def _number(value):
    """API money/rate string → Decimal (a real number in Excel); None stays empty."""
    return None if value in (None, '') else Decimal(str(value))


def _sheet(title):
    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws.sheet_view.rightToLeft = True
    return wb, ws


def _title_block(ws, title, details, columns):
    """Rows 1-2: title and «بندر | دوره | ساخته‌شده»; row 4: header. Returns the first data row."""
    last = get_column_letter(len(columns))
    ws.merge_cells(f'A1:{last}1')
    ws['A1'] = title
    ws['A1'].font = Font(name=FONT, size=14, bold=True, color=BRAND)
    ws.merge_cells(f'A2:{last}2')
    ws['A2'] = '   |   '.join(details)
    ws['A2'].font = Font(name=FONT, size=10, color='4B5563')
    for col, (header, width) in enumerate(columns, start=1):
        cell = ws.cell(row=4, column=col, value=header)
        cell.font = Font(name=FONT, bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor=BRAND)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = Border(bottom=THIN)
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[4].height = 30
    ws.freeze_panes = 'A5'
    return 5


def _write_row(ws, row, values, formats=None, bold=False, fill=None):
    for col, value in enumerate(values, start=1):
        cell = ws.cell(row=row, column=col, value=value)
        cell.font = Font(name=FONT, bold=bold)
        if formats and formats.get(col) and isinstance(value, Decimal):
            cell.number_format = formats[col]
        if fill:
            cell.fill = PatternFill('solid', fgColor=fill)
        cell.border = Border(bottom=THIN)


def xlsx_response(content, filename):
    response = HttpResponse(content, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


def _save(wb):
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def income_statement_workbook(report, branch_label):
    """
    IAS 1 layout: each section's lines and its total, gross profit after COGS, other income
    after the expenses, net profit last.
    """
    wb, ws = _sheet('صورت حساب عاید')
    period = report['period']
    row = _title_block(
        ws,
        'صورت حساب عاید (IAS 1)',
        [f'بندر: {branch_label}',
         f'دوره: {_ltr(period["from"])} تا {_ltr(period["to"])} '
         f'({_ltr(period["from_gregorian"])} – {_ltr(period["to_gregorian"])})',
         f'تعداد معاملات: {report["transactions_count"]}',
         f'ساخته‌شده: {_ltr(report["generated_at"])}'],
        [('کود', 18), ('شرح', 42), ('افغانی', 20), ('دالر', 18)],
    )
    money = {3: MONEY, 4: MONEY}
    totals = report['totals']
    for section in SECTION_TOTAL_KEY:
        lines = [line for line in report['lines'] if line['section'] == section]
        label = lines[0]['section_label'] if lines else dict(StatementLine.Section.choices)[section]
        _write_row(ws, row, [None, label], bold=True, fill=BRAND_LIGHT)
        row += 1
        for line in lines:
            _write_row(ws, row, [line['code'], line['label_fa'], _number(line['amount_afn']),
                                 _number(line['amount_usd'])], money)
            row += 1
        key = SECTION_TOTAL_KEY[section]
        _write_row(ws, row, [None, f'جمع {label}', _number(totals[f'{key}_afn']), _number(totals[f'{key}_usd'])],
                   money, bold=True)
        row += 1
        if section in SUBTOTALS:
            text, key = SUBTOTALS[section]
            _write_row(ws, row, [None, text, _number(totals[f'{key}_afn']), _number(totals[f'{key}_usd'])],
                       money, bold=True, fill=BRAND_LIGHT)
            row += 1
    _write_row(ws, row, [None, 'مفاد / ضرر خالص', _number(totals['net_profit_afn']),
                         _number(totals['net_profit_usd'])], money, bold=True, fill='FDE7D3')
    row += 2
    for warning in report['warnings']:
        ws.cell(row=row, column=1, value=f'⚠ {warning}').font = Font(name=FONT, color='B45309')
        row += 1
    return _save(wb)


TX_INFO_COLUMNS = [
    ('شماره / سریال', 16), ('تاریخ شمسی', 13), ('تاریخ میلادی', 13), ('نوع معامله', 12), ('وضعیت', 14),
    ('بندر', 18), ('تاجر / مشتری', 22), ('شرکت ترانسپورت', 20), ('نمبر پلیت', 14), ('اسم راننده', 16),
    ('نوع ارز', 10), ('نرخ اسعار', 12),
]
TX_RESULT_COLUMNS = [
    ('جمله عواید', 16), ('جمله مصرف', 16), ('جمله عواید (افغانی)', 18), ('جمله مصرف (افغانی)', 18),
    ('مفاد / ضرر (افغانی)', 18), ('مفاد / ضرر (دالر)', 16),
]


def transactions_workbook(transactions, types, branch_label, period_label, generated_at):
    """
    One row per transaction (all filtered rows; usd rates and money must already be attached):
    one column per revenue type, then per expense type (built-in and added, in form order), in
    the row's own currency, then the totals and AFN / USD results, and a totals row of the AFN /
    USD columns. `types` is the operations.money.MoneyTypes the rows were attached with.
    """
    money_types = ([t for t in types.types if t.kind == 'REVENUE']
                   + [t for t in types.types if t.kind != 'REVENUE'])
    suffix = {'REVENUE': 'عاید', 'EXPENSE': 'مصرف'}
    columns = [*TX_INFO_COLUMNS, *[(f'{t.name_fa} ({suffix[t.kind]})', 14) for t in money_types],
               *TX_RESULT_COLUMNS]
    wb, ws = _sheet('معاملات')
    row = _title_block(ws, 'لیست معاملات',
                       [f'بندر: {branch_label}', f'دوره: {period_label}',
                        f'تعداد: {len(transactions)}', f'ساخته‌شده: {_ltr(generated_at)}'], columns)
    first_money = len(TX_INFO_COLUMNS)  # exchange rate; amounts after it
    formats = {col: MONEY for col in range(first_money + 1, len(columns) + 1)}
    formats[first_money] = RATE
    cent = Decimal('0.01')
    sums = {'revenue': Decimal('0'), 'expense': Decimal('0'), 'profit': Decimal('0'), 'profit_usd': Decimal('0')}
    for tx in transactions:
        lines = tx.line_amounts()
        amounts = [getattr(tx, t.column) if t.column else lines.get(t.pk, Decimal('0')) for t in money_types]
        profit_usd = tx.profit_or_loss_usd
        values = [
            tx.serial_number, str(tx.date), str(tx.date.togregorian()), tx.get_trade_type_display(),
            tx.get_status_display(), tx.branch.name, tx.trader.name if tx.trader else None,
            tx.transport_company.name if tx.transport_company else None,
            tx.plate_afghan or tx.plate_foreign or None, tx.driver_name or None,
            tx.currency, tx.exchange_rate, *amounts,
            tx.total_revenue, tx.total_expense, tx.total_revenue_afn.quantize(cent),
            tx.total_expense_afn.quantize(cent), tx.profit_or_loss_afn.quantize(cent),
            None if profit_usd is None else profit_usd.quantize(cent),
        ]
        _write_row(ws, row, values, formats)
        row += 1
        sums['revenue'] += values[-4]
        sums['expense'] += values[-3]
        sums['profit'] += values[-2]
        sums['profit_usd'] = None if sums['profit_usd'] is None or profit_usd is None else sums['profit_usd'] + values[-1]
    total_row = [None] * len(columns)
    total_row[0] = 'جمع'
    total_row[-4:] = [sums['revenue'], sums['expense'], sums['profit'], sums['profit_usd']]
    _write_row(ws, row, total_row, formats, bold=True, fill=BRAND_LIGHT)
    return _save(wb)
