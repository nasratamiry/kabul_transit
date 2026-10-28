"""
Verify the REST API (JWT, role access, port scoping, daily-rate soft lock, FormConfig, Dari digits).

    .venv\\Scripts\\python scripts\\verify_api.py

Runs against the configured database inside one DB transaction that is rolled
back at the end, so no test data is left behind. Test branches use the codes
ZZKBL / ZZTKH so they never clash with real branches.
"""
import os
import sys
from datetime import timedelta
from decimal import Decimal
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'kabulcargo_backend.settings')
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

import django  # noqa: E402

django.setup()

from axes.models import AccessAttempt  # noqa: E402
from django.core.files.uploadedfile import SimpleUploadedFile  # noqa: E402
from django.db import connection, transaction  # noqa: E402
from django.test import TestCase  # noqa: E402
from django.test.utils import CaptureQueriesContext, override_settings  # noqa: E402
from rest_framework.test import APIClient  # noqa: E402
from rest_framework_simplejwt.tokens import AccessToken  # noqa: E402

from accounts.models import Branch, User  # noqa: E402
from config.models import FieldConfig  # noqa: E402
from finance.models import ExchangeRate, Hawala, OfficeExpense  # noqa: E402
from finance.utils import kabul_today  # noqa: E402
from operations.models import Document, Transaction  # noqa: E402
from partners.models import Trader  # noqa: E402

PASSWORD = 'Kabul-Verify-2026!'
failures = []
number = 0


def check(label, ok, detail=''):
    global number
    number += 1
    print(f"  [{'OK' if ok else 'FAIL'}] {number:02d}. {label}{(' — ' + detail) if detail else ''}")
    if not ok:
        failures.append(f'{number:02d}. {label}')


def to_dari(text):
    return str(text).translate(str.maketrans('0123456789', '۰۱۲۳۴۵۶۷۸۹'))


def login(username, password=PASSWORD):
    client = APIClient(HTTP_HOST='127.0.0.1')
    response = client.post('/api/auth/login/', {'username': username, 'password': password}, format='json')
    if response.status_code == 200:
        client.credentials(HTTP_AUTHORIZATION=f"Bearer {response.data['access']}")
    return client, response


def results(response):
    data = response.data
    return data['results'] if isinstance(data, dict) and 'results' in data else data


def counts():
    return (Branch.objects.count(), User.objects.count(), Trader.objects.count(), Transaction.objects.count(),
            Document.objects.count(), ExchangeRate.objects.count(), Hawala.objects.count(),
            OfficeExpense.objects.count(), FieldConfig.objects.count(), AccessAttempt.objects.count())


class Rollback(Exception):
    pass


before = counts()
stored_documents = []
open_responses = []
today = kabul_today()
tomorrow = today + timedelta(days=1)

try:
    with transaction.atomic():
        ExchangeRate.objects.filter(date=today).delete()  # start with no rate today (rolled back)
        branch_a = Branch.objects.create(name='بندر آزمایشی API کابل', code='ZZKBL')
        branch_b = Branch.objects.create(name='بندر آزمایشی API تورخم', code='ZZTKH')
        admin = User.objects.create_superuser('zz_api_admin', None, PASSWORD)
        manager_a = User.objects.create_user('zz_api_mgr_a', None, PASSWORD, role='branch_manager', branch=branch_a)
        manager_b = User.objects.create_user('zz_api_mgr_b', None, PASSWORD, role='branch_manager', branch=branch_b)
        trader = Trader.objects.create(name='تاجر آزمایشی API')
        tx_a = Transaction.objects.create(branch=branch_a, trade_type='EXPORT', serial_number='ZZ-A-1',
                                          trader=trader, plate_afghan='ZZ-111')
        tx_b = Transaction.objects.create(branch=branch_b, trade_type='EXPORT', serial_number='ZZ-B-1',
                                          trader=trader)
        hawala_b = Hawala.objects.create(branch=branch_b, from_party='الف', to_party='ب', amount=100, currency='AFN')

        print('الف) ورود و توکن (Auth)')
        admin_client, response = login(admin.username)
        check('سوپر ادمین وارد می‌شود و access + refresh + user می‌گیرد',
              response.status_code == 200 and {'access', 'refresh', 'user'} <= set(response.data)
              and response.data['user']['role'] == 'super_admin', f'status={response.status_code}')
        admin_refresh = response.data.get('refresh')

        _, response = login(admin.username, 'wrong-password')
        check('رمز اشتباه رد می‌شود (401)', response.status_code == 401, f'status={response.status_code}')

        client_a, response = login(manager_a.username)
        claims = AccessToken(response.data['access'])
        check('توکن نقش و بندر را دارد', claims['role'] == 'branch_manager' and claims['branch_id'] == branch_a.id,
              f"role={claims['role']}, branch_id={claims['branch_id']}")

        response = client_a.get('/api/auth/me/')
        check('/api/auth/me/ کاربر و بندرش را برمی‌گرداند',
              response.status_code == 200 and response.data['username'] == manager_a.username
              and response.data['branch']['id'] == branch_a.id, f'status={response.status_code}')

        response = APIClient(HTTP_HOST='127.0.0.1').get('/api/transactions/')
        check('بدون توکن دسترسی نیست (401)', response.status_code == 401, f'status={response.status_code}')

        anon = APIClient(HTTP_HOST='127.0.0.1')
        refreshed = anon.post('/api/auth/refresh/', {'refresh': admin_refresh}, format='json')
        logout = admin_client.post('/api/auth/logout/', {'refresh': admin_refresh}, format='json')
        after_logout = anon.post('/api/auth/refresh/', {'refresh': admin_refresh}, format='json')
        check('refresh کار می‌کند؛ بعد از خروج همان refresh باطل است',
              refreshed.status_code == 200 and 'access' in refreshed.data and logout.status_code == 200
              and after_logout.status_code == 401,
              f'refresh={refreshed.status_code}, logout={logout.status_code}, after={after_logout.status_code}')
        admin_client, _ = login(admin.username)  # fresh session for the rest
        client_b, _ = login(manager_b.username)

        print('ب) دسترسی نقش‌ها (Role access)')
        response = client_a.get('/api/users/')
        check('مدیر بندر در فهرست کاربران فقط خودش را می‌بیند',
              response.status_code == 200 and [u['id'] for u in results(response)] == [manager_a.id],
              f'status={response.status_code}')

        response = client_a.post('/api/users/', {'username': 'zz_api_hack', 'password': PASSWORD,
                                                 'role': 'super_admin'}, format='json')
        check('مدیر بندر نمی‌تواند کاربر بسازد (403)', response.status_code == 403, f'status={response.status_code}')

        response = admin_client.post('/api/users/', {'username': 'zz_api_new', 'password': PASSWORD,
                                                     'role': 'branch_manager', 'branch': branch_a.id,
                                                     'phone': '۰۷۹۹۱۲۳۴۵۶'}, format='json')
        new_user = User.objects.filter(username='zz_api_new').first()
        check('سوپر ادمین کاربر می‌سازد؛ رمز hash و تلفن انگلیسی می‌شود',
              response.status_code == 201 and new_user is not None and new_user.check_password(PASSWORD)
              and new_user.phone == '0799123456' and 'password' not in response.data,
              f'status={response.status_code}, data={response.data if response.status_code != 201 else ""}')

        denied = client_a.post('/api/branches/', {'name': 'بندر هک', 'code': 'ZZHAK'}, format='json')
        created = admin_client.post('/api/branches/', {'name': 'بندر آزمایشی API هرات', 'code': 'zzhir'},
                                    format='json')
        check('بندر: مدیر 403، سوپر ادمین 201 (کود بزرگ می‌شود)',
              denied.status_code == 403 and created.status_code == 201 and created.data['code'] == 'ZZHIR',
              f'manager={denied.status_code}, admin={created.status_code}')

        branch_url = f"/api/branches/{created.data['id']}/"
        renamed = admin_client.patch(branch_url, {'code': 'ZZNEW'}, format='json')
        same = admin_client.patch(branch_url, {'code': 'ZZHIR', 'name': 'بندر آزمایشی API هرات ۲'}, format='json')
        put = admin_client.put(branch_url, {'name': 'بندر آزمایشی API هرات ۳', 'code': 'ZZPUT'}, format='json')
        check('کود بندر تغییرناپذیر است: PATCH / PUT با کود تازه ۴۰۰ (دری)؛ همان کود + نام تازه ۲۰۰',
              renamed.status_code == 400 and 'code' in renamed.data and put.status_code == 400
              and same.status_code == 200 and Branch.objects.get(pk=created.data['id']).code == 'ZZHIR',
              f'patch={renamed.status_code} same={same.status_code} put={put.status_code}')

        field_config = FieldConfig.objects.filter(branch=branch_a, trade_type='EXPORT').first()
        denied = client_a.patch(f'/api/field-configs/{field_config.id}/', {'sort_order': 1}, format='json')
        allowed = admin_client.patch(f'/api/field-configs/{field_config.id}/', {'sort_order': 1}, format='json')
        check('تنظیم فیلد: مدیر 403، سوپر ادمین 200',
              denied.status_code == 403 and allowed.status_code == 200,
              f'manager={denied.status_code}, admin={allowed.status_code}')

        print('ج) محدودیت بندر (Port scoping)')
        response = client_a.get('/api/transactions/')
        branches_seen = {row['branch'] for row in results(response)}
        check('مدیر A فقط معاملات بندر A را می‌بیند', response.status_code == 200 and branches_seen == {branch_a.id},
              f'branches={branches_seen}')

        response = client_a.get(f'/api/transactions/{tx_b.id}/')
        check('مدیر A معاملهٔ بندر B را باز کرده نمی‌تواند (404)', response.status_code == 404,
              f'status={response.status_code}')

        denied = client_a.patch(f'/api/hawala/{hawala_b.id}/', {'note': 'hack'}, format='json')
        admin_list = admin_client.get('/api/transactions/', {'branch': branch_b.id})
        check('مدیر A حوالهٔ B را تغییر داده نمی‌تواند؛ سوپر ادمین همه را می‌بیند',
              denied.status_code == 404 and admin_list.status_code == 200
              and [row['id'] for row in results(admin_list)] == [tx_b.id],
              f'manager={denied.status_code}, admin={admin_list.status_code}')

        # Soft lock needs today's rate for the spoofing test; add it, then remove for group (د).
        rate = ExchangeRate.objects.create(date=today, currency='USD', rate=Decimal('70.5'))
        response = client_a.post('/api/transactions/', {'branch': branch_b.id, 'trade_type': 'IMPORT',
                                                        'serial_number': 'ZZ-SPOOF'}, format='json')
        spoof = Transaction.objects.filter(serial_number='ZZ-SPOOF').first()
        check('جعل بندر: مدیر A با branch=B می‌فرستد، اما در A ذخیره می‌شود',
              response.status_code == 201 and spoof is not None and spoof.branch_id == branch_a.id
              and spoof.created_by_id == manager_a.id,
              f'status={response.status_code}, saved_branch={spoof.branch_id if spoof else None}')

        moved = client_a.patch(f'/api/transactions/{tx_a.id}/', {'branch': branch_b.id}, format='json')
        tx_a.refresh_from_db()
        upload = client_a.post('/api/documents/', {'doc_type': 'OTHER', 'transaction': tx_b.id, 'title': 'hack',
                                                   # a valid PDF, so the branch check (403) is what refuses it
                                                   'file': SimpleUploadedFile('x.pdf', b'%PDF-1.4\n%%EOF\n')},
                                  format='multipart')
        check('مدیر A معامله را به B برده نمی‌تواند و به معاملهٔ B سند وصل کرده نمی‌تواند',
              moved.status_code == 200 and tx_a.branch_id == branch_a.id and upload.status_code == 403,
              f'patch={moved.status_code}, branch={tx_a.branch_id}, upload={upload.status_code}')
        rate.delete()

        print('د) قفل نرم نرخ روز')
        for label, url, body in (
            ('معامله بدون نرخ امروز → 428', '/api/transactions/',
             {'branch': branch_a.id, 'trade_type': 'EXPORT', 'serial_number': 'ZZ-LOCK'}),
            ('حواله بدون نرخ امروز → 428', '/api/hawala/',
             {'branch': branch_a.id, 'from_party': 'الف', 'to_party': 'ب', 'amount': '10', 'currency': 'USD'}),
            ('مصرف دفتر بدون نرخ امروز → 428', '/api/office-expenses/',
             {'branch': branch_a.id, 'category': 'SALARY', 'amount': '10', 'currency': 'USD'}),
        ):
            response = client_a.post(url, body, format='json')
            check(label, response.status_code == 428 and response.data.get('code') == 'TODAY_RATE_MISSING',
                  f"status={response.status_code}, code={getattr(response, 'data', {}).get('code')}")

        future = client_a.post('/api/exchange-rates/', {'date': str(tomorrow), 'currency': 'USD', 'rate': '71'},
                               format='json')
        response = client_a.post('/api/exchange-rates/', {'date': to_dari(today), 'currency': 'USD',
                                                          'rate': '۷۰٫۵'}, format='json')
        rate = ExchangeRate.objects.filter(date=today, currency='USD').first()
        check('نرخ فردا رد می‌شود (400)؛ نرخ امروز با ارقام دری ثبت می‌شود (201)',
              future.status_code == 400 and response.status_code == 201 and rate is not None
              and rate.rate == Decimal('70.5') and rate.created_by_id == manager_a.id,
              f'future={future.status_code}, today={response.status_code}')

        response = client_a.post('/api/office-expenses/', {'branch': branch_b.id, 'category': 'SALARY',
                                                           'amount': '10', 'currency': 'USD'}, format='json')
        expense = OfficeExpense.objects.filter(pk=response.data.get('id')).first() if response.status_code == 201 else None
        check('بعد از ثبت نرخ، همان درخواست موفق است (نرخ ۷۰.۵ ذخیره، بندر A)',
              expense is not None and expense.rate_used == Decimal('70.5') and expense.branch_id == branch_a.id
              and response.data['amount_afn'] == '705.00',
              f'status={response.status_code}')

        print('ه) محتوای معامله')
        body = {'branch': branch_a.id, 'trade_type': 'TRANSIT', 'serial_number': 'ZZ-T-1', 'trader': trader.id,
                'plate_afghan': 'ZZ-777', 'currency': 'USD', 'exchange_rate': '70.5',
                'customer_payment': '1000', 'khak_puli_cash': '200', 'transit_right': '100'}
        response = client_a.post('/api/transactions/', body, format='json')
        tx = Transaction.objects.filter(serial_number='ZZ-T-1').first()
        check('معامله ثبت شد؛ بیلانس = پرداخت مشتری − مصارف (۷۰۰)',
              response.status_code == 201 and tx is not None and tx.balance == Decimal('700')
              and response.data['balance'] == '700.00' and tx.created_by_id == manager_a.id,
              f'status={response.status_code}, data={response.data if response.status_code != 201 else ""}')

        summary = response.data.get('row_summary', {}) if response.status_code == 201 else {}
        check('row_summary: عاید ۷۰۵۰۰ / ۱۰۰۰ دالر، مصرف ۲۱۱۵۰ افغانی / ۳۰۰ دالر، مفاد ۴۹۳۵۰ افغانی / ۷۰۰ دالر',
              summary == {'total_revenue_afn': '70500.00', 'total_revenue_usd': '1000.00',
                          'total_expense_afn': '21150.00', 'fee_per_dollar_afn': '70.50000000',
                          'total_expense_usd': '300.00', 'profit_or_loss_afn': '49350.00',
                          'profit_or_loss_usd': '700.00'}, str(summary))

        shamsi = client_a.get('/api/transactions/', {'date_from': str(today), 'trade_type': 'TRANSIT'})
        gregorian = client_a.get('/api/transactions/', {'date_to': str(today.togregorian()), 'trade_type': 'TRANSIT'})
        check('تاریخ شمسی برگشت؛ فلتر تاریخ با شمسی و میلادی هر دو کار می‌کند',
              response.data.get('date') == str(today)
              and [r['id'] for r in results(shamsi)] == [tx.id] and [r['id'] for r in results(gregorian)] == [tx.id],
              f"date={response.data.get('date')}, shamsi={shamsi.status_code}, gregorian={gregorian.status_code}")

        response = client_a.get('/api/transactions/', {'search': 'ZZ-777'})
        check('جستجو با نمبر پلیت کار می‌کند', [r['id'] for r in results(response)] == [tx.id],
              f'status={response.status_code}')

        print('و) فورم پویا (FormConfig)')
        response = client_a.get('/api/form-config/', {'branch': branch_b.id, 'trade_type': 'EXPORT'})
        check('مدیر A با branch=B هم فورم بندر A را می‌گیرد',
              response.status_code == 200 and response.data['branch'] == branch_a.id,
              f"status={response.status_code}, branch={response.data.get('branch')}")

        response = admin_client.get('/api/form-config/', {'branch': branch_b.id, 'trade_type': 'EXPORT'})
        missing = admin_client.get('/api/form-config/', {'trade_type': 'EXPORT'})
        check('سوپر ادمین branch=B را می‌گیرد؛ بدون branch → 400',
              response.status_code == 200 and response.data['branch'] == branch_b.id and missing.status_code == 400,
              f'with={response.status_code}, without={missing.status_code}')

        FieldConfig.objects.filter(branch=branch_b, trade_type='EXPORT', field_name='remarks').update(visibility='HIDDEN')
        response = admin_client.get('/api/form-config/', {'branch': branch_b.id, 'trade_type': 'EXPORT'})
        fields = response.data['form_fields']
        orders = [f['sort_order'] for f in fields]
        check('فیلدهای پنهان حذف و بقیه به ترتیب sort_order اند',
              'remarks' not in [f['field_name'] for f in fields] and orders == sorted(orders)
              and all(f['required'] == (f['visibility'] == 'REQUIRED') for f in fields),
              f'{len(fields)} fields')

        print('ز) ارقام دری در API')
        response = client_a.post('/api/transactions/', {'branch': branch_a.id, 'trade_type': 'EXPORT',
                                                        'serial_number': '۴۵۶', 'currency': 'AFN',
                                                        'customer_payment': '۱٬۲۰۰', 'expense_labor': '٢٠٠'},
                                 format='json')
        tx = Transaction.objects.filter(pk=response.data.get('id')).first() if response.status_code == 201 else None
        check('مبلغ و سریال با ارقام دری/عربی ذخیره و انگلیسی می‌شوند',
              tx is not None and tx.serial_number == '456' and tx.customer_payment == Decimal('1200')
              and tx.balance == Decimal('1000'),
              f'status={response.status_code}, data={response.data if response.status_code != 201 else ""}')

        dari = client_a.get('/api/transactions/', {'date': to_dari(today), 'search': '456'})
        bad = client_a.get('/api/transactions/', {'date': 'not-a-date'})
        check('فلتر تاریخ با ارقام دری کار می‌کند؛ تاریخ غلط → 400',
              dari.status_code == 200 and [r['id'] for r in results(dari)] == [tx.id if tx else None]
              and bad.status_code == 400, f'dari={dari.status_code}, bad={bad.status_code}')

        counts_tx = Transaction.objects.count()
        print('ح) باز کردن قفل حساب و حذف محافظت‌شده')
        AccessAttempt.objects.create(username=manager_b.username, ip_address='10.9.9.9', user_agent='verify',
                                     http_accept='', path_info='/api/auth/login/', get_data='', post_data='',
                                     failures_since_start=5)
        listed = {row['username']: row['is_locked'] for row in results(admin_client.get('/api/users/'))}
        locked_out = login(manager_b.username)[1]
        check('حساب قفل‌شده در فهرست کاربران is_locked=true دارد و ورودش ۴۲۳ است',
              listed.get(manager_b.username) is True and listed.get(manager_a.username) is False
              and locked_out.status_code == 423, f'listed={listed.get(manager_b.username)}, login={locked_out.status_code}')

        denied = client_a.post(f'/api/users/{manager_b.id}/unlock/')
        check('مدیر بندر نمی‌تواند قفل را باز کند (403)',
              denied.status_code == 403 and AccessAttempt.objects.filter(username=manager_b.username).exists(),
              f'status={denied.status_code}')

        unlocked = admin_client.post(f'/api/users/{manager_b.id}/unlock/')
        _, relogin = login(manager_b.username)
        check('سوپر ادمین قفل را باز می‌کند: تلاش‌های axes پاک و ورود دوباره ۲۰۰',
              unlocked.status_code == 200 and unlocked.data['reset'] >= 1
              and not AccessAttempt.objects.filter(username=manager_b.username).exists()
              and relogin.status_code == 200, f'status={unlocked.status_code}, data={unlocked.data}, login={relogin.status_code}')

        protected = admin_client.delete(f'/api/branches/{branch_a.id}/')
        check('حذف بندری که کاربر دارد → 409 با پیام دری (نه 500)',
              protected.status_code == 409 and protected.data.get('code') == 'PROTECTED'
              and 'کاربران' in protected.data['detail'] and Branch.objects.filter(pk=branch_a.id).exists(),
              f'status={protected.status_code}, data={getattr(protected, "data", None)}')

        print('ط) پیش‌نمایش خلاصهٔ ردیف و صفحه‌بندی معاملات')
        preview = admin_client.post('/api/transactions/preview-summary/', {
            'branch': branch_a.id, 'currency': 'AFN', 'customer_payment': '۱۰٬۰۰۰',
            'revenue_transport': '2000', 'tariff_expense': '3000', 'expense_labor': 'not-a-number',
        }, format='json')
        usd = admin_client.post('/api/transactions/preview-summary/', {
            'currency': 'USD', 'exchange_rate': '70', 'customer_payment': '100', 'tariff_expense': '30',
        }, format='json')
        check('preview-summary: «عواید ترانسپورت» مصرف حساب می‌شود؛ ورودی خراب = ۰؛ دالری با نرخ خودش',
              preview.status_code == 200 and preview.data['customer_payment_afn'] == '10000.00'
              and preview.data['total_expense_afn'] == '5000.00' and preview.data['profit_or_loss_afn'] == '5000.00'
              and preview.data['profit_or_loss_usd'] == str((Decimal('5000') / Decimal(preview.data['fee_per_dollar_afn']))
                                                           .quantize(Decimal('0.01')))
              and usd.data['profit_or_loss_afn'] == '4900.00' and usd.data['profit_or_loss_usd'] == '70.00'
              and Transaction.objects.count() == counts_tx,
              f'afn={preview.data}, usd={usd.data}')

        other = client_a.post('/api/transactions/preview-summary/', {'branch': branch_b.id, 'customer_payment': '1'},
                              format='json')
        own = client_a.post('/api/transactions/preview-summary/', {'branch': branch_a.id, 'customer_payment': '1'},
                            format='json')
        check('preview-summary: مدیر بندر برای بندر دیگر ۴۰۳، برای بندر خودش ۲۰۰',
              other.status_code == 403 and own.status_code == 200, f'other={other.status_code}, own={own.status_code}')

        paged = admin_client.get('/api/transactions/', {'page_size': 1})
        check('فهرست معاملات صفحه‌بندی دارد (count / next / results)',
              paged.status_code == 200 and isinstance(paged.data, dict) and len(paged.data['results']) == 1
              and paged.data['count'] >= 2 and paged.data['next'], f'data keys={list(getattr(paged, "data", {}) or {})}')

        print('ک) معادل دالری، اسناد، فیلتر، تاریخچه، نرخ مدیر بندر، حذف')
        # 1.1 — an AFN transaction uses the USD rate of ITS date (else the latest earlier one), never today's
        day = today - timedelta(days=20)
        ExchangeRate.objects.filter(currency='USD', date__lt=today).delete()
        ExchangeRate.objects.create(date=day - timedelta(days=2), currency='USD', rate=Decimal('68'))
        ExchangeRate.objects.update_or_create(date=today, currency='USD', defaults={'rate': Decimal('80')})
        old_afn = Transaction.objects.create(branch=branch_a, trade_type='EXPORT', serial_number='ZZ-USD-1',
                                             currency='AFN', date=day, customer_payment=Decimal('6800'),
                                             tariff_expense=Decimal('1360'))
        no_rate = Transaction.objects.create(branch=branch_a, trade_type='EXPORT', serial_number='ZZ-USD-0',
                                             currency='AFN', date=day - timedelta(days=30),
                                             customer_payment=Decimal('100'))
        summary = client_a.get(f'/api/transactions/{old_afn.id}/').data['row_summary']
        none_summary = client_a.get(f'/api/transactions/{no_rate.id}/').data['row_summary']
        check('معاملهٔ افغانی: معادل دالری با نرخ دالرِ تاریخ خودش (۶۸، نه نرخ امروز ۸۰)؛ بی‌نرخ → None',
              summary['fee_per_dollar_afn'] == '68.00000000' and summary['profit_or_loss_usd'] == '80.00'
              and summary['total_expense_usd'] == '20.00' and none_summary['profit_or_loss_usd'] is None,
              f'{summary} | {none_summary}')

        def list_queries(page_size):
            with CaptureQueriesContext(connection) as ctx:
                response = client_a.get('/api/transactions/', {'page_size': page_size, 'date_from': str(day - timedelta(days=40))})
            return len(ctx.captured_queries), len(response.data['results'])
        Transaction.objects.bulk_create([
            Transaction(branch=branch_a, trade_type='IMPORT', serial_number=f'ZZ-Q-{i}', currency='AFN',
                        date=day + timedelta(days=i % 15), customer_payment=Decimal('10'), balance=Decimal('10'))
            for i in range(30)])
        small, small_rows = list_queries(3)
        big, big_rows = list_queries(30)
        check('فهرست معاملات: تعداد queryها با تعداد ردیف‌ها بالا نمی‌رود (بدون N+1 برای نرخ دالر)',
              small == big and big_rows == 30 and small_rows == 3, f'3 rows → {small} queries, 30 rows → {big} queries')

        # 1.2 — documents: authenticated, branch-checked download; no raw /media/ URL
        png = (b'\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06\x00\x00\x00'
               b'\x1f\x15\xc4\x89\x00\x00\x00\rIDATx\x9cc\xf8\x0f\x00\x00\x01\x01\x00\x05\x18\xd8N\x00'
               b'\x00\x00\x00IEND\xaeB`\x82')
        upload = client_a.post('/api/documents/', {'doc_type': 'OTHER', 'transaction': tx_a.id, 'title': 'بارنامه',
                                                    'file': SimpleUploadedFile('بارنامهٔ تورخم.png', png)},
                               format='multipart')
        doc = Document.objects.filter(pk=upload.data.get('id')).first() if upload.status_code == 201 else None
        listed = client_a.get('/api/documents/', {'transaction': tx_a.id})
        body = str(upload.data) + str(listed.data)
        check('سند آپلود شد؛ پاسخ API هیچ نشانی خام /media/ ندارد، فقط download_url',
              doc is not None and '/media/' not in body and upload.data['download_url'] == f'/api/documents/{doc.id}/download/'
              and upload.data['content_type'] == 'image/png' and 'file' not in upload.data,
              f'status={upload.status_code}, keys={sorted(upload.data)}')
        if doc:
            stored_documents.append(doc.file.name)
            own = client_a.get(f'/api/documents/{doc.id}/download/')
            content = b''.join(own.streaming_content) if own.status_code == 200 else b''
            attach = client_a.get(f'/api/documents/{doc.id}/download/', {'download': '1'})
            other = client_b.get(f'/api/documents/{doc.id}/download/')
            anon = APIClient(HTTP_HOST='127.0.0.1').get(f'/api/documents/{doc.id}/download/')
            # A FileResponse keeps its file open until closed (Windows cannot delete it then), but
            # response.close() fires request_finished, which closes the DB connection: close them
            # only after the rolled-back transaction (see `finally`).
            open_responses.extend([own, attach])
            check('دانلود: بندر خودش ۲۰۰ (همان بایت‌ها، نوع درست، نام دری)؛ بندر دیگر ۴۰۴؛ بدون توکن ۴۰۱',
                  own.status_code == 200 and content == png and own['Content-Type'] == 'image/png'
                  and "filename*=utf-8''" in own['Content-Disposition'] and own['Content-Disposition'].startswith('inline')
                  and attach['Content-Disposition'].startswith('attachment')
                  and other.status_code == 404 and anon.status_code == 401,
                  f"own={own.status_code} {own.get('Content-Disposition')}, other={other.status_code}, anon={anon.status_code}")
            with override_settings(DOCUMENT_X_ACCEL_PREFIX='/protected-media'):
                accel = client_a.get(f'/api/documents/{doc.id}/download/')
            check('production: X-Accel-Redirect به محل internal nginx (Django فایل را نمی‌فرستد)',
                  accel.status_code == 200 and accel['X-Accel-Redirect'].startswith('/protected-media/documents/')
                  and not getattr(accel, 'streaming', False) and accel.content == b'',
                  accel.get('X-Accel-Redirect', ''))

        # 1.3 — has_documents filter
        with_docs = [r['id'] for r in results(client_a.get('/api/transactions/', {'has_documents': 'true', 'page_size': 200, 'date_from': str(day - timedelta(days=60))}))]
        without = [r['id'] for r in results(client_a.get('/api/transactions/', {'has_documents': 'false', 'page_size': 200, 'date_from': str(day - timedelta(days=60))}))]
        check('فیلتر has_documents: true فقط معاملهٔ سنددار، false بقیه',
              with_docs == [tx_a.id] and tx_a.id not in without and old_afn.id in without, f'true={with_docs}, false={len(without)}')

        # 1.4 — history endpoint (changed fields only), branch-scoped
        client_a.patch(f'/api/transactions/{old_afn.id}/', {'customer_payment': '7000', 'branch': branch_a.id}, format='json')
        history = client_a.get(f'/api/transactions/{old_afn.id}/history/')
        newest = history.data[0] if history.status_code == 200 and history.data else {}
        trader_history = admin_client.get(f'/api/traders/{trader.id}/history/')
        hidden = client_b.get(f'/api/transactions/{old_afn.id}/history/')
        check('تاریخچهٔ معامله: فقط فیلد تغییرکرده (پرداخت مشتری 6800 → 7000) با کاربر و تاریخ دوگانه؛ بندر دیگر ۴۰۴؛ تاجر هم دارد',
              history.status_code == 200 and len(history.data) == 2 and newest.get('user') == manager_a.username
              and newest.get('changes') == [{'field': 'customer_payment', 'label': 'پرداخت مشتری', 'old': '6800.00', 'new': '7000.00'}]
              and history.data[-1]['changes'] == [] and ' | ' in newest.get('date', '')
              and hidden.status_code == 404 and trader_history.status_code == 200,
              f"{history.status_code} {newest} | other={hidden.status_code} trader={trader_history.status_code}")

        # 1.5 — branch manager: today's rate only; no past / future, no edit, no delete
        today_eur = client_a.post('/api/exchange-rates/', {'date': str(today), 'currency': 'EUR', 'rate': '82'}, format='json')
        no_date = client_a.post('/api/exchange-rates/', {'currency': 'PKR', 'rate': '0.25'}, format='json')
        past = client_a.post('/api/exchange-rates/', {'date': str(today - timedelta(days=1)), 'currency': 'IRR', 'rate': '0.0016'}, format='json')
        rate_id = today_eur.data.get('id') if today_eur.status_code == 201 else None
        patch = client_a.patch(f'/api/exchange-rates/{rate_id}/', {'rate': '90'}, format='json')
        delete = client_a.delete(f'/api/exchange-rates/{rate_id}/')
        admin_past = admin_client.post('/api/exchange-rates/', {'date': str(today - timedelta(days=1)), 'currency': 'IRR', 'rate': '0.0016'}, format='json')
        admin_delete = admin_client.delete(f'/api/exchange-rates/{rate_id}/')
        check('نرخ: مدیر بندر نرخ امروز می‌سازد (۲۰۱)؛ تاریخ گذشته / ویرایش / حذف ۴۰۳؛ سوپر ادمین هر سه را می‌تواند',
              today_eur.status_code == 201 and no_date.status_code == 201 and past.status_code == 403
              and patch.status_code == 403 and delete.status_code == 403
              and admin_past.status_code == 201 and admin_delete.status_code == 204,
              f'today={today_eur.status_code}, no-date={no_date.status_code}, past={past.status_code}, patch={patch.status_code}, '
              f'delete={delete.status_code}, admin past={admin_past.status_code}, admin delete={admin_delete.status_code}')

        # 1.6 — who may delete a transaction (documented, unchanged): manager own branch yes, other branch no
        throwaway = Transaction.objects.create(branch=branch_a, trade_type='IMPORT', serial_number='ZZ-DEL-1')
        other_branch = client_b.delete(f'/api/transactions/{throwaway.id}/')
        own_branch = client_a.delete(f'/api/transactions/{throwaway.id}/')
        check('حذف معامله: مدیر بندر معاملهٔ بندر خودش را حذف می‌کند (۲۰۴)، بندر دیگر ۴۰۴',
              other_branch.status_code == 404 and own_branch.status_code == 204, f'other={other_branch.status_code}, own={own_branch.status_code}')

        print('ل) حواله و مصارف دفتر: ارقام دری، بازهٔ تاریخ، جمع‌ها، snapshot نرخ، بندر، ۴۲۸، تاریخچه')
        from finance.reports import income_statement  # noqa: E402
        from kabulcargo_backend.utils import normalize_digits, parse_decimal  # noqa: E402
        check('normalize_digits / parse_decimal: «٫» → «.» و «٬» حذف (۱٬۲۳۴٫۵۰ → 1234.50)',
              normalize_digits('۱۲٫۵') == '12.5' and parse_decimal('۱٬۲۳۴٫۵۰') == Decimal('1234.50')
              and parse_decimal('٣٬٠٠٠٫٢٥') == Decimal('3000.25'), normalize_digits('۱۲٫۵'))

        start, end = today - timedelta(days=4), today
        ExchangeRate.objects.update_or_create(date=start, currency='USD', defaults={'rate': Decimal('80')})
        ExchangeRate.objects.update_or_create(date=start, currency='PKR', defaults={'rate': Decimal('0.25')})
        Hawala.objects.filter(branch=branch_a).delete()
        OfficeExpense.objects.filter(branch=branch_a).delete()
        h_afn = client_a.post('/api/hawala/', {'branch': branch_a.id, 'date': str(start), 'from_party': 'الف', 'to_party': 'ب',
                                                'amount': '۱٬۲۵۰٫۷۵', 'currency': 'AFN', 'fee': '۵۰'}, format='json')
        h_usd = client_a.post('/api/hawala/', {'branch': branch_a.id, 'date': str(start), 'from_party': 'ج', 'to_party': 'د',
                                                'amount': '100', 'currency': 'USD', 'fee': '2'}, format='json')
        h_old = Hawala.objects.create(branch=branch_a, date=start - timedelta(days=30), from_party='x', to_party='y',
                                      amount=Decimal('5'), currency='AFN', fee=Decimal('1'))
        e_rent = client_a.post('/api/office-expenses/', {'branch': branch_a.id, 'date': str(start), 'category': 'RENT',
                                                          'amount': '۳٬۰۰۰', 'currency': 'AFN'}, format='json')
        e_pkr = client_a.post('/api/office-expenses/', {'branch': branch_a.id, 'date': str(start), 'category': 'FUEL',
                                                         'amount': '400', 'currency': 'PKR'}, format='json')
        e_salary = client_a.post('/api/office-expenses/', {'branch': branch_a.id, 'date': str(end), 'category': 'SALARY',
                                                            'amount': '۲۵۰', 'currency': 'USD'}, format='json')
        check('حواله / مصرف دفتر با ارقام دری و «٫» ذخیره می‌شود؛ rate_used = نرخ همان روز',
              h_afn.status_code == 201 and h_afn.data['amount'] == '1250.75' and h_afn.data['rate_used'] is None
              and h_usd.status_code == 201 and h_usd.data['rate_used'] == '80.00000000' and h_usd.data['fee_afn'] == '160.00'
              and e_rent.data.get('amount') == '3000.00' and e_pkr.data.get('rate_used') == '0.25000000'
              and e_pkr.data.get('amount_afn') == '100.00' and e_salary.status_code == 201,
              f'{h_afn.status_code} {h_afn.data.get("amount")} | {h_usd.data.get("rate_used")} | {e_pkr.data.get("amount_afn")}')

        period = {'date_from': to_dari(start), 'date_to': str(end.togregorian()), 'page_size': 1}
        hl = client_a.get('/api/hawala/', period)
        ol = client_a.get('/api/office-expenses/', period)
        check('بازهٔ تاریخ (شمسی دری / میلادی) + صفحه‌بندی + totals برای کل فیلتر (نه فقط صفحه)',
              hl.status_code == 200 and hl.data['count'] == 2 and len(hl.data['results']) == 1
              and h_old.id not in [r['id'] for r in hl.data['results']]
              and hl.data['totals']['amount']['by_currency'] == {'AFN': '1250.75', 'USD': '100.00'}
              and hl.data['totals']['fee']['afn'] == '210.00' and hl.data['totals']['fee']['usd'] == '2.63'
              and ol.data['count'] == 3 and ol.data['totals']['afn'] == '23100.00'
              and ol.data['totals']['salary']['afn'] == '20000.00' and ol.data['totals']['non_salary']['afn'] == '3100.00'
              and ol.data['totals']['groups']['FUEL']['afn'] == '100.00',
              f"hawala={hl.data.get('count')} {hl.data.get('totals')} | office={ol.data.get('count')} {ol.data.get('totals')}")

        report = income_statement(branch_id=branch_a.id, date_from=start, date_to=end)
        lines = {line['code']: line for line in report['lines']}
        check('جمع صفحه‌ها = سطرهای گزارش عاید (OPEX_OFFICE / OPEX_SALARY / OPEX_HAWALA_FEE، افغانی و دالر)',
              lines['OPEX_OFFICE']['amount_afn'] == ol.data['totals']['non_salary']['afn']
              and lines['OPEX_SALARY']['amount_afn'] == ol.data['totals']['salary']['afn']
              and lines['OPEX_HAWALA_FEE']['amount_afn'] == hl.data['totals']['fee']['afn']
              and lines['OPEX_OFFICE']['amount_usd'] == ol.data['totals']['non_salary']['usd']
              and lines['OPEX_HAWALA_FEE']['amount_usd'] == hl.data['totals']['fee']['usd'],
              f"report={[(c, lines[c]['amount_afn']) for c in ('OPEX_OFFICE', 'OPEX_SALARY', 'OPEX_HAWALA_FEE')]}")

        ExchangeRate.objects.filter(date=start, currency='USD').update(rate=Decimal('95'))
        again = client_a.get('/api/hawala/', period)
        stored = client_a.get(f"/api/hawala/{h_usd.data['id']}/")
        check('تغییر بعدی نرخ روز، رکورد قبلی را عوض نمی‌کند (snapshot)',
              stored.data['rate_used'] == '80.00000000' and again.data['totals']['fee']['afn'] == '210.00',
              f"rate_used={stored.data['rate_used']}, fee afn={again.data['totals']['fee']['afn']}")

        spoof = client_a.post('/api/office-expenses/', {'branch': branch_b.id, 'date': str(today), 'category': 'OTHER',
                                                         'amount': '1', 'currency': 'AFN'}, format='json')
        b_sees = [r['id'] for r in results(client_b.get('/api/office-expenses/', {'page_size': 200}))]
        check('مدیر بندر برای بندر دیگر ثبت نمی‌کند (در بندر خودش ذخیره می‌شود) و بندر دیگر آن را نمی‌بیند',
              spoof.status_code == 201 and spoof.data['branch'] == branch_a.id and spoof.data['id'] not in b_sees,
              f"branch={spoof.data.get('branch')}")

        ExchangeRate.objects.filter(date=today, currency='USD').delete()
        blocked_h = client_a.post('/api/hawala/', {'branch': branch_a.id, 'from_party': 'a', 'to_party': 'b', 'amount': '1',
                                                   'currency': 'AFN'}, format='json')
        blocked_e = client_a.post('/api/office-expenses/', {'branch': branch_a.id, 'category': 'OTHER', 'amount': '1',
                                                            'currency': 'AFN'}, format='json')
        client_a.patch(f"/api/hawala/{h_afn.data['id']}/", {'fee': '60', 'branch': branch_a.id}, format='json')
        history = client_a.get(f"/api/hawala/{h_afn.data['id']}/history/")
        other_history = client_b.get(f"/api/office-expenses/{e_rent.data['id']}/history/")
        check('بدون نرخ امروز: ثبت حواله و مصرف دفتر ۴۲۸؛ تاریخچه فقط فیلد تغییرکرده؛ بندر دیگر ۴۰۴',
              blocked_h.status_code == 428 and blocked_e.status_code == 428
              and history.data[0]['changes'] == [{'field': 'fee', 'label': 'مصرف حواله', 'old': '50.00', 'new': '60.00'}]
              and other_history.status_code == 404,
              f'{blocked_h.status_code}/{blocked_e.status_code}, {history.data[0] if history.status_code == 200 else history.status_code}')

        print('م) داشبورد و خروجی Excel')
        from io import BytesIO  # noqa: E402

        from django.db import connection  # noqa: E402
        from django.test.utils import CaptureQueriesContext  # noqa: E402
        from openpyxl import load_workbook  # noqa: E402

        from finance.dashboard import previous_shamsi_month  # noqa: E402
        from finance.reports import current_shamsi_month  # noqa: E402

        month_from, month_to = current_shamsi_month()
        prev_from, prev_to = previous_shamsi_month(month_from)
        Transaction.objects.create(branch=branch_a, trade_type='IMPORT', serial_number='ZZ-DASH-1', date=month_from,
                                   currency='AFN', customer_payment=Decimal('9000'), tariff_expense=Decimal('1500'))
        Transaction.objects.create(branch=branch_a, trade_type='IMPORT', serial_number='ZZ-DASH-0', date=prev_from,
                                   currency='AFN', customer_payment=Decimal('4000'), tariff_expense=Decimal('500'))
        with CaptureQueriesContext(connection) as queries:
            dash = client_a.get('/api/dashboard/summary/', {'branch': branch_b.id})
        expected = income_statement(branch_id=branch_a.id, date_from=month_from, date_to=month_to)['summary']
        expected_prev = income_statement(branch_id=branch_a.id, date_from=prev_from, date_to=prev_to)['summary']
        current = {k: v for k, v in dash.data['current'].items() if k != 'period'}
        previous = {k: v for k, v in dash.data['previous'].items() if k != 'period'}
        daily_net = sum(Decimal(d['net_profit_afn']) for d in dash.data['daily'])
        check(f'داشبورد مدیر بندر = گزارش عاید همان ماه و ماه قبل (بندر خودش، {len(queries)} کوئری)',
              dash.status_code == 200 and current == expected and previous == expected_prev
              and str(daily_net.quantize(Decimal('0.01'))) == expected['net_profit_afn']
              and 'branches' not in dash.data and dash.data['branch'] == branch_a.id
              and any('بندر خود' in w for w in dash.data['warnings']) and len(queries) <= 25,
              f"current={current} expected={expected} daily={daily_net} queries={len(queries)}")

        admin_dash = admin_client.get('/api/dashboard/summary/')
        rows = {row['branch']['id']: row for row in admin_dash.data.get('branches', [])}
        row_a = {k: v for k, v in rows.get(branch_a.id, {}).items() if k != 'branch'}
        top = admin_client.get('/api/dashboard/summary/', {'branch': branch_a.id}).data
        check('سوپر ادمین: ردیف مقایسهٔ هر بندر = گزارش همان بندر؛ با branch بدون مقایسه؛ تاجران برتر و آخرین معاملات',
              admin_dash.status_code == 200 and row_a == expected and 'branches' not in top
              and len(top['recent_transactions']) <= 10 and top['recent_transactions'][0]['serial_number']
              and isinstance(top['top_traders'], list) and 'is_set' in top['today_rate'],
              f"row_a={row_a} recent={len(top['recent_transactions'])}")

        export = client_a.get('/api/income-statement/export/',
                              {'date_from': str(month_from), 'date_to': str(month_to), 'branch': branch_b.id})
        report = client_a.get('/api/income-statement/', {'date_from': str(month_from), 'date_to': str(month_to)}).data
        ws = load_workbook(BytesIO(export.content)).active
        cells = {ws.cell(row=r, column=2).value: ws.cell(row=r, column=3).value for r in range(5, ws.max_row + 1)}
        codes = [ws.cell(row=r, column=1).value for r in range(5, ws.max_row + 1)
                 if ws.cell(row=r, column=1).value and not str(ws.cell(row=r, column=1).value).startswith('⚠')]
        net = cells.get('مفاد / ضرر خالص')
        check('خروجی Excel گزارش عاید: RTL، عدد واقعی، سطرها و مفاد خالص = API، بندر خود مدیر',
              export.status_code == 200 and export['Content-Type'].startswith('application/vnd.openxmlformats')
              and ws.sheet_view.rightToLeft and ws.freeze_panes == 'A5'
              and isinstance(net, Decimal | float | int) and f'{net:.2f}' == report['totals']['net_profit_afn']
              and codes == [line['code'] for line in report['lines']] and branch_a.name in ws['A2'].value,
              f"{export.status_code} net={net!r} vs {report['totals']['net_profit_afn']} codes={len(codes)}")

        tx_export = client_a.get('/api/transactions/export/')
        listed = results(client_a.get('/api/transactions/', {'page_size': 200}))
        foreign = load_workbook(BytesIO(client_a.get('/api/transactions/export/', {'branch': branch_b.id}).content)).active
        ws = load_workbook(BytesIO(tx_export.content)).active
        data_rows = [r for r in ws.iter_rows(min_row=5, values_only=True)]
        body, total = data_rows[:-1], data_rows[-1]
        api_profit = sum(Decimal(r['row_summary']['profit_or_loss_afn']) for r in listed)
        check('خروجی Excel معاملات: همهٔ ردیف‌های فیلترشده (نه یک صفحه)، فقط بندر خود، جمع مفاد = API',
              tx_export.status_code == 200 and len(body) == len(listed)
              and {r[0] for r in body} == {r['serial_number'] for r in listed}
              and total[0] == 'جمع' and f'{total[-2]:.2f}' == f'{api_profit:.2f}'
              and all(isinstance(r[12], Decimal | float | int) for r in body)
              and foreign.max_row == 5,  # another branch's id: empty, like the list (title, header, totals row)
              f"rows={len(body)} api={len(listed)} profit={total[-2]} vs {api_profit}")

        tariff_rows = results(client_a.get('/api/transactions/', {'nonzero': 'tariff_expense', 'page_size': 200}))
        bad_field = client_a.get('/api/transactions/', {'nonzero': 'branch'})
        OfficeExpense.objects.create(branch=branch_a, date=today, category='SALARY', amount=Decimal('5'), currency='AFN')
        OfficeExpense.objects.create(branch=branch_a, date=today, category='RENT', amount=Decimal('7'), currency='AFN')
        salary = {r['category'] for r in results(client_a.get('/api/office-expenses/', {'kind': 'salary'}))}
        others = {r['category'] for r in results(client_a.get('/api/office-expenses/', {'kind': 'non_salary'}))}
        check('فیلترهای drill-down: nonzero=<فیلد> فقط ردیف‌های غیر صفر (فیلد ناشناخته ۴۰۰)؛ kind=salary / non_salary',
              tariff_rows and all(Decimal(r['tariff_expense']) != 0 for r in tariff_rows)
              and len(tariff_rows) < len(listed) and bad_field.status_code == 400
              and salary == {'SALARY'} and others and 'SALARY' not in others,
              f'tariff={len(tariff_rows)}/{len(listed)} bad={bad_field.status_code} salary={salary} others={others}')

        # Deleting a document (or its transaction) removes the stored file too, after the commit.
        from django.core.files.storage import default_storage  # noqa: E402
        gone = client_a.post('/api/documents/', {'doc_type': 'OTHER', 'transaction': tx_a.id, 'file': SimpleUploadedFile('حذفی.png', png)},
                             format='multipart')
        gone_doc = Document.objects.get(pk=gone.data['id'])
        stored_documents.append(gone_doc.file.name)
        on_disk = default_storage.exists(gone_doc.file.name)
        with TestCase.captureOnCommitCallbacks(execute=True):
            deleted = client_a.delete(f'/api/documents/{gone_doc.id}/')
        check('حذف سند فایل آن را هم از دیسک پاک می‌کند (بعد از commit)',
              on_disk and deleted.status_code == 204 and not default_storage.exists(gone_doc.file.name),
              f'before={on_disk}, status={deleted.status_code}, after={default_storage.exists(gone_doc.file.name)}')

        anonymous = APIClient(HTTP_HOST='127.0.0.1', REMOTE_ADDR='10.9.9.9')
        probes = [anonymous.get('/api/health/') for _ in range(25)]  # more than the anon throttle (20/min)
        body = probes[-1].json()
        check('GET /api/health/ بدون ورود ۲۰۰، throttle نمی‌شود، فقط status و checks (بدون نسخه یا خطا)',
              all(r.status_code == 200 for r in probes) and set(body) == {'status', 'checks'}
              and body == {'status': 'ok', 'checks': {'database': True, 'cache': True}}
              and 'no-store' in probes[-1]['Cache-Control'],
              f'statuses={sorted({r.status_code for r in probes})} body={body}')

        raise Rollback
except Rollback:
    pass
finally:
    # Uploaded files live on disk, outside the rolled-back transaction.
    from django.core.files.storage import default_storage  # noqa: E402
    for response in open_responses:
        response.close()
    for name in stored_documents:
        default_storage.delete(name)

print('ی) پاک‌کاری')
check('همهٔ داده‌های آزمایشی rollback شد', counts() == before, f'before={before}, after={counts()}')

print()
if failures:
    print(f'{len(failures)} از {number} تست ناکام:')
    for label in failures:
        print(f'  - {label}')
    sys.exit(1)
print(f'همهٔ {number} تست موفق بود.')
